from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)


def set_run_font(run, size=12, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(
    text,
    *,
    size=12,
    bold=False,
    italic=False,
    align="justify",
    space_after=8,
    first_line=True,
):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(space_after)
    if first_line and align == "justify":
        pf.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_custom(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(12)
    pf.space_after = Pt(12)
    pf.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True)
    return p


add_para(
    "CENTRO UNIVERSITÁRIO CATÓLICO SALESIANO AUXILIUM",
    align="center",
    bold=True,
    first_line=False,
    space_after=0,
)
add_para("UniSALESIANO", align="center", bold=True, first_line=False, space_after=24)
add_para(
    "MIGUEL RODRIGUES LINO",
    align="center",
    bold=True,
    first_line=False,
    space_after=48,
)
add_para(
    "LEVANTAMENTO BIBLIOGRÁFICO",
    align="center",
    bold=True,
    first_line=False,
    space_after=6,
)
add_para(
    "Tema: StudyHub — Plataforma de Organização de Estudos",
    align="center",
    italic=True,
    first_line=False,
    space_after=48,
)
add_para("Araçatuba", align="center", first_line=False, space_after=0)
add_para("2026", align="center", first_line=False, space_after=24)

doc.add_page_break()

add_heading_custom("1 INTRODUÇÃO")
add_para(
    "O presente levantamento bibliográfico tem como objetivo reunir fontes científicas "
    "relacionadas ao tema do Trabalho de Conclusão de Curso (TCC) intitulado StudyHub — "
    "Plataforma de Organização de Estudos. O StudyHub é um sistema web voltado a apoiar "
    "estudantes na organização de disciplinas, tarefas, sessões de estudo, cadernos digitais, "
    "acompanhamento de desempenho e recursos de assistência tecnológica à aprendizagem."
)
add_para(
    "A pesquisa bibliográfica concentrou-se em temas correlatos à proposta do sistema, "
    "especialmente: autorregulação da aprendizagem, gerenciamento de tempo acadêmico, "
    "uso de tecnologias digitais (aplicativos, diários de aprendizagem e inteligência artificial) "
    "para apoiar o estudo de universitários."
)

add_heading_custom("2 PROCEDIMENTOS DE BUSCA")
add_para(
    "As buscas foram realizadas nas plataformas comumente indicadas para pesquisa acadêmica "
    "em cursos de graduação (bases de acesso aberto e portais científicos), a saber:"
)

for item in [
    "Google Acadêmico (Google Scholar) — https://scholar.google.com/",
    "SciELO — https://www.scielo.br/",
    "Portal de Periódicos CAPES — https://www.periodicos.capes.gov.br/",
    "BDTD / IBICT (Biblioteca Digital Brasileira de Teses e Dissertações) — https://bdtd.ibict.br/",
    "PePSIC / bases correlatas de Psicologia — https://pepsic.bvsalud.org/",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("• " + item)
    set_run_font(run)

add_para(
    'Foram utilizados descritores em português e em inglês, combinados com operadores booleanos, '
    'tais como: "autorregulação da aprendizagem", "gerenciamento de tempo", '
    '"estudantes universitários", "aplicativo", "tecnologia educacional", '
    '"self-regulated learning", "time management" e "mobile learning diaries".'
)
add_para(
    'Observação: a pasta "Material de Apoio" (Unidade 1 – Aula 4) não estava disponível '
    "localmente no momento da elaboração deste arquivo. Caso o material institucional indique "
    "plataformas adicionais específicas do UniSALESIANO, recomenda-se complementar este "
    "levantamento com busca também nessas bases."
)

add_heading_custom("3 FONTES SELECIONADAS")
add_para(
    "Foram selecionadas cinco fontes, conforme solicitação da atividade: dois artigos em língua "
    "portuguesa, um artigo em língua inglesa, um livro em língua portuguesa e uma dissertação "
    "de mestrado em língua portuguesa."
)

add_heading_custom("3.1 Artigos em língua portuguesa")

add_para("Artigo 1", bold=True, first_line=False, space_after=6)
add_para(
    "SILVA, Maria Antônia Romão da; ALLIPRANDINI, Paula Mariza Zedu. Sistemas de diários de "
    "aprendizagem para apoiar a autorregulação da aprendizagem no ensino superior: uma revisão "
    "sistemática de literatura. Revista Brasileira de Estudos Pedagógicos, Brasília, v. 105, "
    "p. e6121, 2024. DOI: https://doi.org/10.24109/2176-6681.rbep.105.6121. Disponível em: "
    "https://www.scielo.br/j/rbeped/a/dj3GfZhbBgkDVzwBxVj7xSH/. Acesso em: 14 out. 2026."
)
add_para(
    "Relação com o TCC: o artigo discute sistemas digitais (incluindo aplicativos e recursos "
    "com inteligência artificial) para apoiar a autorregulação da aprendizagem no ensino superior, "
    "tema alinhado às funcionalidades do StudyHub (planejamento, monitoramento do estudo e "
    "suporte tecnológico ao aluno)."
)
add_para(
    "Base consultada: SciELO / Google Acadêmico.",
    italic=True,
    first_line=False,
)

add_para("Artigo 2", bold=True, first_line=False, space_after=6)
add_para(
    "GANDA, Danielle Ribeiro; BORUCHOVITCH, Evely. Intervenção em autorregulação da aprendizagem "
    "com alunos do ensino superior: análise da produção científica. Estudos Interdisciplinares "
    "em Psicologia, Londrina, v. 10, n. 3, p. 3-25, 2019. DOI: "
    "https://doi.org/10.5433/2236-6407.2019v10n3p03. Disponível em: "
    "https://ojs.uel.br/revistas/uel/index.php/eip/article/view/38752. Acesso em: 14 out. 2026."
)
add_para(
    "Relação com o TCC: a revisão evidencia o impacto positivo de programas de intervenção em "
    "autorregulação no ensino superior, fundamentando a relevância de ferramentas que ajudem "
    "o estudante a aprender a organizar e gerenciar o próprio estudo."
)
add_para(
    "Base consultada: Google Acadêmico / periódico de acesso aberto.",
    italic=True,
    first_line=False,
)

add_heading_custom("3.2 Artigo em língua inglesa")
add_para(
    "BROADBENT, Jaclyn; PANADERO, Ernesto; FULLER-TYSZKIEWICZ, Matthew. Effects of mobile-app "
    "learning diaries vs online training on specific self-regulated learning components. "
    "Educational Technology Research and Development, [s. l.], v. 68, p. 2351-2372, 2020. "
    "DOI: https://doi.org/10.1007/s11423-020-09781-6."
)
add_para(
    "Relação com o TCC: o estudo demonstra efeitos positivos de diários de aprendizagem via "
    "aplicativo móvel e treinamento online sobre componentes da autorregulação (incluindo "
    "estratégias metacognitivas e gerenciamento de tempo), reforçando a utilidade de sistemas "
    "digitais de apoio ao estudo, como o StudyHub."
)
add_para(
    "Base consultada: Google Scholar / SpringerLink.",
    italic=True,
    first_line=False,
)

add_heading_custom("3.3 Livro em língua portuguesa")
add_para(
    "FRISON, Lourdes Maria Bragagnolo; BORUCHOVITCH, Evely (org.). Autorregulação da aprendizagem: "
    "cenários, desafios, perspectivas para o contexto educativo. Petrópolis: Vozes, 2020."
)
add_para(
    "Relação com o TCC: a obra reúne fundamentos teóricos e discussões práticas sobre "
    "autorregulação da aprendizagem em diferentes níveis de ensino, inclusive Ensino Superior, "
    "servindo de base conceitual para o desenvolvimento de uma plataforma de organização de estudos."
)
add_para(
    "Base consultada: Google Acadêmico / catálogos editoriais.",
    italic=True,
    first_line=False,
)

add_heading_custom("3.4 Dissertação de mestrado em língua portuguesa")
add_para(
    "MARQUES, Bruna Guimarães. Promoção do gerenciamento de tempo em universitários: contribuições "
    "para a vida acadêmica. 2019. 51 f. Dissertação (Mestrado em Psicologia) — Instituto de Educação, "
    "Universidade Federal Rural do Rio de Janeiro, Seropédica, 2019. Disponível em: "
    "https://rima.ufrrj.br/jspui/handle/20.500.14407/14525. Acesso em: 14 out. 2026."
)
add_para(
    "Relação com o TCC: a dissertação trata do gerenciamento de tempo em universitários "
    "(estabelecimento de metas, cumprimento de prazos, planejamento e organização), competências "
    "diretamente relacionadas aos módulos de tarefas, sessões e calendário do StudyHub."
)
add_para(
    "Base consultada: BDTD / repositório institucional (UFRRJ).",
    italic=True,
    first_line=False,
)

add_heading_custom("4 CONSIDERAÇÕES FINAIS")
add_para(
    "As fontes selecionadas convergem para a importância da autorregulação da aprendizagem e do "
    "gerenciamento de tempo no desempenho acadêmico, bem como para o potencial de tecnologias "
    "digitais (aplicativos, diários eletrônicos e recursos de apoio inteligente) no suporte ao "
    "estudante. Esse conjunto teórico e empírico fundamenta a proposta do StudyHub como plataforma "
    "de organização de estudos."
)

add_heading_custom("REFERÊNCIAS")
refs = [
    (
        "BROADBENT, Jaclyn; PANADERO, Ernesto; FULLER-TYSZKIEWICZ, Matthew. Effects of mobile-app "
        "learning diaries vs online training on specific self-regulated learning components. "
        "Educational Technology Research and Development, [s. l.], v. 68, p. 2351-2372, 2020. "
        "DOI: https://doi.org/10.1007/s11423-020-09781-6."
    ),
    (
        "FRISON, Lourdes Maria Bragagnolo; BORUCHOVITCH, Evely (org.). Autorregulação da aprendizagem: "
        "cenários, desafios, perspectivas para o contexto educativo. Petrópolis: Vozes, 2020."
    ),
    (
        "GANDA, Danielle Ribeiro; BORUCHOVITCH, Evely. Intervenção em autorregulação da aprendizagem "
        "com alunos do ensino superior: análise da produção científica. Estudos Interdisciplinares "
        "em Psicologia, Londrina, v. 10, n. 3, p. 3-25, 2019. DOI: "
        "https://doi.org/10.5433/2236-6407.2019v10n3p03. Disponível em: "
        "https://ojs.uel.br/revistas/uel/index.php/eip/article/view/38752. Acesso em: 14 out. 2026."
    ),
    (
        "MARQUES, Bruna Guimarães. Promoção do gerenciamento de tempo em universitários: contribuições "
        "para a vida acadêmica. 2019. 51 f. Dissertação (Mestrado em Psicologia) — Instituto de Educação, "
        "Universidade Federal Rural do Rio de Janeiro, Seropédica, 2019. Disponível em: "
        "https://rima.ufrrj.br/jspui/handle/20.500.14407/14525. Acesso em: 14 out. 2026."
    ),
    (
        "SILVA, Maria Antônia Romão da; ALLIPRANDINI, Paula Mariza Zedu. Sistemas de diários de "
        "aprendizagem para apoiar a autorregulação da aprendizagem no ensino superior: uma revisão "
        "sistemática de literatura. Revista Brasileira de Estudos Pedagógicos, Brasília, v. 105, "
        "p. e6121, 2024. DOI: https://doi.org/10.24109/2176-6681.rbep.105.6121. Disponível em: "
        "https://www.scielo.br/j/rbeped/a/dj3GfZhbBgkDVzwBxVj7xSH/. Acesso em: 14 out. 2026."
    ),
]

for r in refs:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(-1.25)
    pf.left_indent = Cm(1.25)
    pf.space_after = Pt(12)
    run = p.add_run(r)
    set_run_font(run)

out = (
    Path(__file__).resolve().parent
    / "Levantamento_Bibliografico_StudyHub_UniSALESIANO.docx"
)
doc.save(out)
print(f"saved {out}")
